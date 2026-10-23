"""Build the evaluation corpus.

For every base workbook:
  old.xlsx        base, recalculated and saved by LibreOffice
  new.xlsx        all edits applied, recalculated
  single-<i>.xlsx edit i alone, recalculated (the counterfactual oracle for causes)
  stale.xlsx      value edits written into old.xlsx's XML without recalculating
                  (what Excel in manual-calc mode, or a script, leaves behind)
  uncached.xlsx   edits applied by openpyxl, which writes formulas without values
manifest.json lists each edit with its address in the new file.
"""
import json
import os
import random
import re
import shutil
import sys
import zipfile

import openpyxl
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

import lo
import models

OUT = 'cases'


def addr(r0, c0):
    return f'{get_column_letter(c0 + 1)}{r0 + 1}'


def apply(doc, e):
    kind, sh, a, v = e
    if kind == 'value':
        lo.set_value(doc, sh, a, v)
    elif kind == 'formula':
        lo.set_formula(doc, sh, a, v)
    elif kind == 'insert':
        lo.insert_rows(doc, sh, a, 1)
        if v:
            for j, x in enumerate(v):
                lo.set_value(doc, sh, addr(a, j), x)


def shifted(e, inserts):
    """Address of a value/formula edit after the inserts (inserts are applied last)."""
    kind, sh, a, v = e
    if kind == 'insert':
        return None
    r, c = coordinate_to_tuple(a)
    for (_, ish, row0, _) in inserts:
        if ish == sh and r - 1 >= row0:
            r += 1
    return f'{get_column_letter(c)}{r}'


def stale_copy(src, dst, edits):
    """Rewrite cached <v> of edited input cells in place, leaving formulas' stale values."""
    zin = zipfile.ZipFile(src)
    wb = openpyxl.load_workbook(src, read_only=True)
    sheet_file = {}
    rels = zin.read('xl/_rels/workbook.xml.rels').decode()
    wbx = zin.read('xl/workbook.xml').decode()
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        tgt = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % m.group(2), rels) or re.search(r'Target="([^"]+)"[^>]*Id="%s"' % m.group(2), rels)
        sheet_file[m.group(1)] = 'xl/' + tgt.group(1).lstrip('/').replace('xl/', '')
    patches = {}
    for (kind, sh, a, v) in edits:
        if kind != 'value' or not isinstance(v, (int, float)):
            continue
        patches.setdefault(sheet_file[sh], []).append((a, v))
    zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in patches:
            x = data.decode()
            for a, v in patches[item.filename]:
                pat = re.compile(r'(<c r="%s"[^>]*?)(?: t="[^"]*")?([^>]*>)(?:<v>[^<]*</v>|<is>.*?</is>)?(</c>)' % a)
                if pat.search(x):
                    x = pat.sub(lambda m: f'{m.group(1)}{m.group(2)}<v>{v}</v>{m.group(3)}', x, count=1)
                else:
                    print('  stale: no cell', a, file=sys.stderr)
            data = x.encode()
        zout.writestr(item, data)
    zout.close()


def uncached_copy(src, dst, edits):
    wb = openpyxl.load_workbook(src)
    for (kind, sh, a, v) in edits:
        if kind == 'value':
            wb[sh][a] = v
        elif kind == 'formula':
            wb[sh][a] = v
        elif kind == 'insert':
            wb[sh].insert_rows(a + 1)  # openpyxl does NOT adjust formulas: a known trap
            for j, x in enumerate(v or []):
                wb[sh].cell(a + 1, j + 1, x)
    wb.save(dst)


def build(name, base, edits, singles=True):
    d = os.path.join(OUT, name)
    os.makedirs(d, exist_ok=True)
    doc = lo.open_doc(base)
    lo.recalc_save(doc, os.path.join(d, 'old.xlsx'))
    lo.close(doc)
    old = os.path.join(d, 'old.xlsx')
    plain = [e for e in edits if e[0] != 'insert']
    inserts = [e for e in edits if e[0] == 'insert']
    doc = lo.open_doc(old)
    for e in plain + inserts:
        apply(doc, e)
    lo.recalc_save(doc, os.path.join(d, 'new.xlsx'))
    lo.close(doc)
    manifest = {'name': name, 'edits': []}
    for i, e in enumerate(edits):
        manifest['edits'].append({'kind': e[0], 'sheet': e[1], 'address': e[2] if e[0] != 'insert' else None,
                                  'row0': e[2] if e[0] == 'insert' else None, 'newAddress': shifted(e, inserts),
                                  'single': f'single-{i}.xlsx' if (singles and e[0] != 'insert') else None})
        if singles and e[0] != 'insert':
            doc = lo.open_doc(old)
            apply(doc, e)
            lo.recalc_save(doc, os.path.join(d, f'single-{i}.xlsx'))
            lo.close(doc)
    # Leave-one-out ("but-for") oracle: all edits except i.
    if singles and len(plain) > 1:
        for i, e in enumerate(edits):
            if e[0] == 'insert':
                continue
            doc = lo.open_doc(old)
            for j, f in enumerate(plain + inserts):
                if f is not e:
                    apply(doc, f)
            lo.recalc_save(doc, os.path.join(d, f'loo-{i}.xlsx'))
            lo.close(doc)
            manifest['edits'][i]['loo'] = f'loo-{i}.xlsx'
    if any(e[0] == 'value' for e in edits) and not inserts:
        stale_copy(old, os.path.join(d, 'stale.xlsx'), edits)
        manifest['stale'] = 'stale.xlsx'
    uncached_copy(old, os.path.join(d, 'uncached.xlsx'), edits)
    manifest['uncached'] = 'uncached.xlsx'
    json.dump(manifest, open(os.path.join(d, 'manifest.json'), 'w'), indent=1)
    print('built', name, len(edits), 'edits')


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    os.makedirs('bases', exist_ok=True)
    which = sys.argv[1:] or list(models.MODELS)
    for name in which:
        if name in models.MODELS:
            path = f'bases/{name}.xlsx'
            edits = models.MODELS[name](path)
            build(name, path, edits)
    lo.stop()
