"""Realistic workbook models, written with openpyxl (formulas only) and then
recalculated and saved by LibreOffice, so cached values are genuine.

Each model also lists the edits a reviewer might make, as
(kind, sheet, address, payload):
  ('value', sheet, 'B3', 0.07)         edit an input
  ('formula', sheet, 'C9', '=...')     edit a formula
  ('insert', sheet, row0, None)        insert a row before 0-based row0
"""
import random

import openpyxl
from openpyxl.workbook.defined_name import DefinedName


def _name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def three_statement(path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = 'Assumptions'
    rows = [('Revenue growth', 0.08), ('Gross margin', 0.42), ('Opex % revenue', 0.21),
            ('Tax rate', 0.25), ('Capex % revenue', 0.05), ('Depreciation years', 7),
            ('DSO days', 45), ('Base revenue', 120000)]
    for i, (k, v) in enumerate(rows, start=2):
        a[f'A{i}'] = k
        a[f'B{i}'] = v
    a['A1'], a['B1'] = 'Driver', 'Value'
    _name(wb, 'Growth', "Assumptions!$B$2")
    _name(wb, 'TaxRate', "Assumptions!$B$5")
    years = list(range(2025, 2035))
    cols = [openpyxl.utils.get_column_letter(c) for c in range(2, 2 + len(years))]
    inc = wb.create_sheet('Income')
    inc['A1'] = 'Line'
    for c, y in zip(cols, years):
        inc[f'{c}1'] = y
    lines = ['Revenue', 'COGS', 'Gross profit', 'Opex', 'EBITDA', 'Depreciation', 'EBIT', 'Tax', 'Net income']
    for i, l in enumerate(lines, start=2):
        inc[f'A{i}'] = l
    for j, c in enumerate(cols):
        prev = cols[j - 1] if j else None
        inc[f'{c}2'] = '=Assumptions!$B$9' if j == 0 else f'={prev}2*(1+Growth)'
        inc[f'{c}3'] = f'=-{c}2*(1-Assumptions!$B$3)'
        inc[f'{c}4'] = f'={c}2+{c}3'
        inc[f'{c}5'] = f'=-{c}2*Assumptions!$B$4'
        inc[f'{c}6'] = f'={c}4+{c}5'
        inc[f'{c}7'] = f'=-Cash!{c}3/Assumptions!$B$7' if j == 0 else f'=-SUM(Cash!$B$3:{c}3)/Assumptions!$B$7'
        inc[f'{c}8'] = f'={c}6+{c}7'
        inc[f'{c}9'] = f'=-MAX(0,{c}8)*TaxRate'
        inc[f'{c}10'] = f'={c}8+{c}9'
    cash = wb.create_sheet('Cash')
    cash['A1'] = 'Line'
    for c, y in zip(cols, years):
        cash[f'{c}1'] = y
    for i, l in enumerate(['Net income', 'Capex', 'Receivables', 'Change in AR', 'Free cash flow', 'Cumulative FCF'], start=2):
        cash[f'A{i}'] = l
    for j, c in enumerate(cols):
        prev = cols[j - 1] if j else None
        cash[f'{c}2'] = f'=Income!{c}10'
        cash[f'{c}3'] = f'=Income!{c}2*Assumptions!$B$6'
        cash[f'{c}4'] = f'=Income!{c}2*Assumptions!$B$8/365'
        cash[f'{c}5'] = f'={c}4' if j == 0 else f'={c}4-{prev}4'
        cash[f'{c}6'] = f'={c}2-{c}3-{c}5-Income!{c}7'
        cash[f'{c}7'] = f'={c}6' if j == 0 else f'={prev}7+{c}6'
    summ = wb.create_sheet('Summary')
    summ['A1'], summ['B1'] = 'Metric', 'Value'
    for i, (k, f) in enumerate([('Total revenue', '=SUM(Income!B2:K2)'), ('Total net income', '=SUM(Income!B10:K10)'),
                                 ('Avg net margin', '=B3/B2'), ('Ending cash', '=Cash!K7'),
                                 ('Payback year', '=IFERROR(INDEX(Cash!B1:K1,MATCH(TRUE,INDEX(Cash!B7:K7>0,0),0)),"never")'),
                                 ('Margin flag', '=IF(B4>0.15,"healthy",IF(B4>0.05,"thin","loss"))')], start=2):
        summ[f'A{i}'] = k
        summ[f'B{i}'] = f
    wb.save(path)
    edits = [('value', 'Assumptions', 'B2', 0.10), ('value', 'Assumptions', 'B5', 0.27),
             ('formula', 'Income', 'F9', '=-MAX(0,F8)*TaxRate*1.1'), ('value', 'Assumptions', 'B8', 50)]
    return edits


def budget(path, n=600, seed=3):
    r = random.Random(seed)
    wb = openpyxl.Workbook()
    t = wb.active
    t.title = 'Transactions'
    t.append(['id', 'month', 'category', 'amount', 'approved'])
    cats = ['Travel', 'Software', 'Payroll', 'Marketing', 'Office', 'Legal']
    for i in range(n):
        t.append([1000 + i, r.randint(1, 12), r.choice(cats), round(r.uniform(20, 5000), 2), r.choice(['yes', 'yes', 'yes', 'no'])])
    s = wb.create_sheet('Summary')
    s['A1'] = 'category'
    for m in range(1, 13):
        s.cell(1, m + 1, f'M{m}')
    s.cell(1, 14, 'Total')
    s.cell(1, 15, 'Budget')
    s.cell(1, 16, 'Variance')
    s.cell(1, 17, 'Status')
    last = n + 1
    for i, c in enumerate(cats, start=2):
        s.cell(i, 1, c)
        for m in range(1, 13):
            col = openpyxl.utils.get_column_letter(m + 1)
            s[f'{col}{i}'] = f'=SUMIFS(Transactions!$D$2:$D${last},Transactions!$C$2:$C${last},$A{i},Transactions!$B$2:$B${last},{m},Transactions!$E$2:$E${last},"yes")'
        s[f'N{i}'] = f'=SUM(B{i}:M{i})'
        s[f'O{i}'] = r.choice([40000, 60000, 80000, 120000])
        s[f'P{i}'] = f'=N{i}-O{i}'
        s[f'Q{i}'] = f'=IF(P{i}>0,"over",IF(P{i}>-0.1*O{i},"near","under"))'
    tot = len(cats) + 2
    s[f'A{tot}'] = 'Total'
    for col in [openpyxl.utils.get_column_letter(c) for c in range(2, 17)]:
        s[f'{col}{tot}'] = f'=SUM({col}2:{col}{tot - 1})'
    wb.save(path)
    edits = [('value', 'Transactions', 'D40', 9999.0), ('value', 'Transactions', 'E77', 'no'),
             ('value', 'Summary', 'O4', 70000), ('insert', 'Transactions', 300, [2001, 6, 'Travel', 1234.5, 'yes'])]
    return edits


def inventory(path, products=200, orders=2500, seed=5):
    r = random.Random(seed)
    wb = openpyxl.Workbook()
    p = wb.active
    p.title = 'Products'
    p.append(['sku', 'name', 'price', 'cost', 'margin'])
    for i in range(products):
        price = round(r.uniform(5, 300), 2)
        p.append([f'SKU-{i:04d}', f'Item {i}', price, round(price * r.uniform(0.4, 0.8), 2), None])
        p[f'E{i + 2}'] = f'=(C{i + 2}-D{i + 2})/C{i + 2}'
    o = wb.create_sheet('Orders')
    o.append(['order', 'sku', 'qty', 'unit price', 'line total', 'cost', 'profit'])
    for i in range(orders):
        k = i + 2
        o.append([50000 + i, f'SKU-{r.randrange(products):04d}', r.randint(1, 20)])
        o[f'D{k}'] = f'=VLOOKUP(B{k},Products!$A$2:$D${products + 1},3,FALSE)'
        o[f'E{k}'] = f'=C{k}*D{k}'
        o[f'F{k}'] = f'=C{k}*INDEX(Products!$D$2:$D${products + 1},MATCH(B{k},Products!$A$2:$A${products + 1},0))'
        o[f'G{k}'] = f'=E{k}-F{k}'
    s = wb.create_sheet('Report')
    s['A1'], s['B1'] = 'metric', 'value'
    last = orders + 1
    s['A2'], s['B2'] = 'Revenue', f'=SUM(Orders!E2:E{last})'
    s['A3'], s['B3'] = 'Profit', f'=SUM(Orders!G2:G{last})'
    s['A4'], s['B4'] = 'Margin', '=B3/B2'
    s['A5'], s['B5'] = 'Top SKU units', f'=MAX(Orders!C2:C{last})'
    s['A6'], s['B6'] = 'SKU-0007 revenue', f'=SUMIF(Orders!B2:B{last},"SKU-0007",Orders!E2:E{last})'
    s['A7'], s['B7'] = 'Orders > 1000', f'=COUNTIF(Orders!E2:E{last},">1000")'
    wb.save(path)
    edits = [('value', 'Products', 'C9', 99.0), ('value', 'Orders', 'C100', 50),
             ('formula', 'Report', 'B4', '=ROUND(B3/B2,4)'), ('value', 'Products', 'D150', 1.0)]
    return edits


def loan(path, months=360):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Loan'
    s['A1'], s['B1'] = 'Principal', 350000
    s['A2'], s['B2'] = 'Annual rate', 0.065
    s['A3'], s['B3'] = 'Years', 30
    s['A4'], s['B4'] = 'Payment', '=-PMT(B2/12,B3*12,B1)'
    s['A5'], s['B5'] = 'Extra payment', 0
    s['A7'], s['B7'], s['C7'], s['D7'], s['E7'] = 'month', 'opening', 'interest', 'principal', 'closing'
    for m in range(1, months + 1):
        k = m + 7
        s[f'A{k}'] = m
        s[f'B{k}'] = '=B1' if m == 1 else f'=E{k - 1}'
        s[f'C{k}'] = f'=ROUND(B{k}*$B$2/12,2)'
        s[f'D{k}'] = f'=MIN(B{k},$B$4-C{k}+$B$5)'
        s[f'E{k}'] = f'=MAX(0,B{k}-D{k})'
    s['G1'], s['H1'] = 'Total interest', f'=SUM(C8:C{months + 7})'
    s['G2'], s['H2'] = 'Months to payoff', f'=COUNTIF(B8:B{months + 7},">0")'
    wb.save(path)
    return [('value', 'Loan', 'B2', 0.07), ('value', 'Loan', 'B5', 200)]


def commission(path, reps=150, seed=9):
    r = random.Random(seed)
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Reps'
    s.append(['rep', 'region', 'quota', 'sales', 'attain', 'tier', 'rate', 'commission', 'capped'])
    for i in range(reps):
        k = i + 2
        s.append([f'R{i:03d}', r.choice(['N', 'S', 'E', 'W']), r.choice([200000, 300000, 400000]), round(r.uniform(50000, 600000), 0)])
        s[f'E{k}'] = f'=D{k}/C{k}'
        s[f'F{k}'] = f'=IF(E{k}>=1.2,"gold",IF(E{k}>=1,"silver",IF(E{k}>=0.7,"bronze","none")))'
        s[f'G{k}'] = f'=IF(F{k}="gold",Rates!$B$2,IF(F{k}="silver",Rates!$B$3,IF(F{k}="bronze",Rates!$B$4,0)))'
        s[f'H{k}'] = f'=ROUND(D{k}*G{k},2)'
        s[f'I{k}'] = f'=MIN(H{k},Rates!$B$6)'
    t = wb.create_sheet('Rates')
    for i, (k, v) in enumerate([('gold', 0.1), ('silver', 0.07), ('bronze', 0.04), ('none', 0), ('cap', 45000)], start=2):
        t[f'A{i}'], t[f'B{i}'] = k, v
    t['D1'], t['E1'] = 'region', 'payout'
    for i, reg in enumerate(['N', 'S', 'E', 'W'], start=2):
        t[f'D{i}'] = reg
        t[f'E{i}'] = f'=SUMIF(Reps!$B$2:$B${reps + 1},D{i},Reps!$I$2:$I${reps + 1})'
    t['D7'], t['E7'] = 'total', '=SUM(E2:E5)'
    wb.save(path)
    # A sales edit that crosses a tier threshold (an IF branch flips), a rate edit, a quota edit.
    return [('value', 'Reps', 'D10', 610000.0), ('value', 'Rates', 'B3', 0.08), ('value', 'Reps', 'C20', 100000)]


def chain(path, n):
    """A long running-total chain: the worst case for dependency walks."""
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Chain'
    s.append(['i', 'x', 'running', 'share'])
    for i in range(n):
        k = i + 2
        s.append([i, (i * 7919) % 1000])
        s[f'C{k}'] = '=B2' if i == 0 else f'=C{k - 1}+B{k}'
        s[f'D{k}'] = f'=C{k}/$F$1'
    s['F1'] = f'=C{n + 1}'
    wb.save(path)
    return [('value', 'Chain', 'B3', 5000)]


def lookup_grid(path, n):
    """Many independent rows (n), each with a few formulas: wide-but-shallow."""
    r = random.Random(11)
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Data'
    s.append(['id', 'a', 'b', 'c', 'sum', 'ratio', 'flag'])
    for i in range(n):
        k = i + 2
        s.append([i, r.randint(1, 100), r.randint(1, 100), r.randint(1, 100)])
        s[f'E{k}'] = f'=B{k}+C{k}+D{k}'
        s[f'F{k}'] = f'=B{k}/E{k}'
        s[f'G{k}'] = f'=IF(F{k}>$J$1,"high","low")'
    s['I1'], s['J1'] = 'threshold', 0.4
    s['I2'], s['J2'] = 'total', f'=SUM(E2:E{n + 1})'
    s['I3'], s['J3'] = 'highs', f'=COUNTIF(G2:G{n + 1},"high")'
    wb.save(path)
    edits = [('value', 'Data', 'J1', 0.35)] + [('value', 'Data', f'B{r.randint(2, n + 1)}', 77) for _ in range(20)]
    return edits


MODELS = {
    'three_statement': three_statement,
    'budget': budget,
    'inventory': inventory,
    'loan': loan,
    'commission': commission,
}
