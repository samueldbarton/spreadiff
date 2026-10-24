"""Differential test of spreadiff's formula engine: write probe formulas into
a workbook, have LibreOffice recalculate it, and save it with the results.
`compare.ts` then evaluates every probe with spreadiff's engine and compares.

Usage: python3 build.py [count] [seed]   (from this folder; writes work/probe.xlsx)
"""
import os
import random
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'trace-eval'))
import lo  # noqa: E402

COUNT = int(sys.argv[1]) if len(sys.argv) > 1 else 1500
SEED = int(sys.argv[2]) if len(sys.argv) > 2 else 7
WORK = os.path.join(os.path.dirname(__file__), 'work')
os.makedirs(WORK, exist_ok=True)

# ---- Inputs ---------------------------------------------------------------
NUMS = [0, 1, -1, 2.5, -2.5, 0.1, 1e10, 3, 7, 0.333333333333333, 100, 45292, -0.5, 1.5, 12345.678]
TEXT = ['abc', 'ABC', ' x ', '10', '1e3', 'TRUE', 'a*b', 'a?c', 'Ab C', '-5', 'apple pie', 'Zed', '', '3.5', 'a~b']
MIXED = [True, False, None, 5, '5', None, True, 0, 'x', -3, 2, None, 'TRUE', 1, 4]
SORTED = [10, 20, 30, 40, 50, 60, 70, 80, 90, 100]
NAMES = ['a', 'b', 'c', 'd', 'e', 'f', 'g', 'h', 'i', 'j']
UNSORTED = [50, 10, 90, 30, 70, 20, 100, 40, 80, 60]

# Number formats: TEXT(value, code) for common and finance codes (the grid shows cells through the same formatter).
FORMAT_VALUES = [0, 1234.5678, -1234.5678, 0.0789, 12345678, 0.5, 45292.75, -0.004, 1.5, 999.995]
FORMAT_CODES = [
    'General', '0', '0.00', '#,##0', '#,##0.00', '0%', '0.0%', '0.00E+00', '##0.0E+0', '# ?/?', '# ??/??', '?/8',
    '$#,##0_);($#,##0)', '$#,##0.00_);[Red]($#,##0.00)', '#,##0_);(#,##0)', '#,##0.0,,"m"', '#,##0,"k"', '0.0"x"',
    '_(* #,##0_);_(* (#,##0);_(* "-"_);_(@_)', '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)',
    '[$€-407] #,##0.00', '[>=1000]#,##0,"k";0', '0;-0;"zero"', '#,##0.00;(#,##0.00);"-"', '000-0000', '#.##', '0.0#',
    'd-mmm-yy', 'mmmm d, yyyy', 'dddd', 'h:mm AM/PM', '[h]:mm:ss', 'mm:ss.0', 'yyyy-mm-dd hh:mm', '[$-409]mmm-yy',
]


def text_probe(v, code):
    return '=TEXT(%r,"%s")' % (v, code.replace('"', '""'))


CURATED = [
    # Operators: precedence, unary minus, percent, concatenation, comparison across types
    '=-2^2', '=2^-1', '=-(2^2)', '=2^3^2', '=10%', '=5%%', '=1+2*3-4/5', '=(1+2)*3', '=2*-3',
    '=1&2', '="a"&TRUE', '=1/3&""', '=0.1+0.2&""', '=1E+20&""', '=-0&""', '=123456789012345678&""',
    '="a"<"B"', '="a"="A"', '="10"=10', '=TRUE>1', '="x">99', '=FALSE<TRUE', '=""=0', '=In!C3=0', '=In!C3=""',
    '=In!C3+1', '=In!C3&"x"', '=TRUE+TRUE', '="3"+"4"', '="3"*2', '=" 3 "+1', '="1e3"+0', '="abc"+1',
    '=1/0', '=0/0', '=SQRT(-1)', '=#N/A+1', '=IFERROR(1/0,"e")', '=ISERROR(#REF!)', '=NA()',
    # Rounding and number functions
    '=ROUND(2.5,0)', '=ROUND(-2.5,0)', '=ROUND(1.005,2)', '=ROUND(1234.5678,-2)', '=ROUNDUP(-1.2,0)', '=ROUNDDOWN(-1.9,0)',
    '=INT(-1.5)', '=TRUNC(-1.5)', '=MOD(-7,3)', '=MOD(7,-3)', '=MOD(7.5,2)', '=MOD(5,0)', '=MROUND(10,3)', '=MROUND(-10,-3)',
    '=CEILING(2.1,1)', '=CEILING(-2.1,-1)', '=FLOOR(2.9,1)', '=FLOOR(-2.1,-1)', '=CEILING.MATH(-2.1)', '=FLOOR.MATH(-2.1)',
    '=EVEN(1.1)', '=EVEN(-1.1)', '=ODD(2)', '=ODD(-2)', '=FACT(5)', '=FACT(0)', '=POWER(2,0.5)', '=POWER(0,0)', '=SIGN(-0.1)',
    '=LOG(8,2)', '=LOG(100)', '=LN(EXP(1))', '=ATAN2(1,1)', '=ABS(-3)', '=SQRT(16)',
    # Aggregates: text, booleans and blanks in ranges vs arguments
    '=SUM(In!C1:C15)', '=SUM(TRUE,1)', '=SUM("3",1)', '=AVERAGE(In!C1:C15)', '=COUNT(In!C1:C15)', '=COUNTA(In!C1:C15)',
    '=COUNTBLANK(In!C1:C15)', '=MAX(In!B1:B15)', '=MIN(In!C1:C15)', '=MAX()', '=PRODUCT(In!C1:C15)', '=MEDIAN(In!A1:A15)',
    '=STDEV(In!A1:A15)', '=STDEVP(In!A1:A15)', '=VAR(In!A1:A15)', '=VARP(In!A1:A15)', '=LARGE(In!A1:A15,2)', '=SMALL(In!A1:A15,3)',
    '=RANK(7,In!A1:A15)', '=RANK(7,In!A1:A15,1)', '=SUMSQ(In!A1:A5)', '=SUMPRODUCT(In!E1:E10,In!G1:G10)',
    '=SUMPRODUCT((In!E1:E10>40)*In!G1:G10)', '=SUBTOTAL(9,In!A1:A15)', '=SUBTOTAL(1,In!E1:E10)', '=AVERAGE(In!B1:B15)',
    # Criteria
    '=COUNTIF(In!B1:B15,"a*")', '=COUNTIF(In!B1:B15,"a?c")', '=COUNTIF(In!B1:B15,"a~*b")', '=COUNTIF(In!B1:B15,"10")',
    '=COUNTIF(In!B1:B15,10)', '=COUNTIF(In!A1:A15,">2")', '=COUNTIF(In!A1:A15,"<>0")', '=COUNTIF(In!C1:C15,"")',
    '=COUNTIF(In!C1:C15,"<>")', '=COUNTIF(In!C1:C15,TRUE)', '=COUNTIF(In!C1:C15,"TRUE")', '=COUNTIF(In!B1:B15,"abc")',
    '=COUNTIF(In!A1:A15,">=1E+10")', '=SUMIF(In!E1:E10,">50",In!G1:G10)', '=SUMIF(In!F1:F10,"c",In!E1:E10)',
    '=AVERAGEIF(In!E1:E10,"<45")', '=SUMIFS(In!E1:E10,In!G1:G10,">40",In!F1:F10,"<>d")', '=COUNTIFS(In!E1:E10,">20",In!E1:E10,"<80")',
    '=MAXIFS(In!E1:E10,In!G1:G10,"<60")', '=MINIFS(In!G1:G10,In!E1:E10,">=50")', '=AVERAGEIFS(In!G1:G10,In!E1:E10,">0")',
    '=COUNTIF(In!B1:B15,"*")', '=COUNTIF(In!A1:A15,"=0")', '=COUNTIF(In!B1:B15,"<b")', '=SUMIF(In!A1:A15,"<0")',
    # Lookups
    '=VLOOKUP(35,In!E1:F10,2)', '=VLOOKUP(35,In!E1:F10,2,FALSE)', '=VLOOKUP(5,In!E1:F10,2)', '=VLOOKUP(100,In!E1:F10,2,TRUE)',
    '=VLOOKUP("C",In!F1:F10,1,FALSE)', '=VLOOKUP(30,In!E1:F10,3,FALSE)', '=HLOOKUP(20,In!E12:N13,2,FALSE)',
    '=MATCH(45,In!E1:E10)', '=MATCH(45,In!E1:E10,1)', '=MATCH(50,In!G1:G10,0)', '=MATCH("d",In!F1:F10,0)', '=MATCH("D*",In!F1:F10,0)',
    '=MATCH(45,In!E1:E10,-1)', '=MATCH(5,In!E1:E10,1)', '=INDEX(In!E1:F10,3,2)', '=INDEX(In!E1:E10,0)', '=INDEX(In!E1:F10,11,1)',
    '=XLOOKUP(70,In!G1:G10,In!F1:F10)', '=XLOOKUP(65,In!E1:E10,In!F1:F10,"none")', '=XLOOKUP(65,In!E1:E10,In!F1:F10,,-1)',
    '=XLOOKUP(65,In!E1:E10,In!F1:F10,,1)', '=LOOKUP(55,In!E1:E10,In!F1:F10)', '=CHOOSE(2,"a","b","c")', '=CHOOSE(4,1,2,3)',
    '=ROWS(In!E1:F10)', '=COLUMNS(In!E1:F10)', '=ROW(In!E5)', '=COLUMN(In!F1)', '=INDIRECT("In!E3")', '=OFFSET(In!E1,2,1)',
    '=SUM(OFFSET(In!E1,1,0,3,1))',
    # Logic
    '=IF(0,"y","n")', '=IF("",1,2)', '=IF(In!C3,1,2)', '=AND(TRUE,1,"x")', '=AND(In!C1:C3)', '=OR(0,0)', '=XOR(TRUE,TRUE,TRUE)',
    '=NOT(0)', '=IFS(1>2,"a",2>1,"b")', '=IFS(FALSE,1)', '=SWITCH(2,1,"a",2,"b","z")', '=SWITCH(9,1,"a")', '=IFNA(NA(),"n")',
    '=IF(1,,2)', '=IF(0,1)',
    # Text
    '=LEN(1/3)', '=LEN(In!C3)', '=LEFT("abc",5)', '=LEFT("abc")', '=RIGHT("abc",0)', '=MID("abcdef",3,2)', '=MID("abc",5,2)',
    '=MID("abc",0,1)', '=UPPER("aBc")', '=PROPER("hello WORLD o\'neil")', '=TRIM("  a   b  ")', '=CLEAN("a"&CHAR(9)&"b")',
    '=REPT("ab",3)', '=EXACT("a","A")', '=SUBSTITUTE("aaa","a","b",2)', '=SUBSTITUTE("aaa","a","b")', '=REPLACE("abcdef",2,3,"X")',
    '=FIND("b","abcb")', '=FIND("B","abc")', '=SEARCH("B","abc")', '=SEARCH("?c","abc")', '=SEARCH("x","abc")', '=FIND("","abc")',
    '=VALUE("1,234")', '=VALUE("12%")', '=VALUE(" 3 ")', '=VALUE("abc")', '=VALUE("1e3")', '=CHAR(65)', '=CODE("A")', '=CODE("")',
    '=CONCATENATE("a",1,TRUE)', '=CONCAT(In!F1:F3)', '=TEXTJOIN("-",TRUE,"a","","b")', '=TEXTJOIN("-",FALSE,"a","","b")',
    '=T(1)', '=T("x")', '=N("5")', '=N(TRUE)', '=FIXED(1234.567,1)', '=FIXED(1234.567,-1,TRUE)',
    '=TEXT(1234.5,"0.00")', '=TEXT(0.25,"0%")', '=TEXT(1234.5,"#,##0")', '=TEXT(45292,"yyyy-mm-dd")', '=TEXT(-1.5,"0.0")',
    '=TEXT(0.5,"0.0%")', '=TEXT(3,"000")',
    # Dates
    '=DATE(2024,2,30)', '=DATE(2024,13,1)', '=DATE(2024,0,1)', '=DATE(99,1,1)', '=YEAR(45292)', '=MONTH(45292)', '=DAY(45292)',
    '=WEEKDAY(45292)', '=WEEKDAY(45292,2)', '=WEEKDAY(45292,3)', '=EDATE(45322,1)', '=EOMONTH(45292,1)', '=EOMONTH(45292,-1)',
    '=DAYS(45322,45292)', '=TIME(25,0,0)', '=HOUR(0.75)', '=MINUTE(0.5104)', '=SECOND(0.50001)', '=DATEVALUE("2024-01-31")',
    '=DATE(1900,3,1)-DATE(1900,2,28)', '=YEAR(60)', '=DAY(59)',
    # Finance
    '=PMT(0.05/12,360,200000)', '=PMT(0,12,1200)', '=FV(0.05,10,-100)', '=PV(0.05,10,-100)', '=NPER(0.05,-100,1000)',
    '=IPMT(0.05/12,1,360,200000)', '=PPMT(0.05/12,1,360,200000)', '=NPV(0.1,-100,50,60,70)', '=IRR(In!H1:H5)', '=RATE(10,-100,800)',
    '=PMT(0.05,10,1000,0,1)',
    # Info and conversions
    '=ISBLANK(In!C3)', '=ISBLANK("")', '=ISNUMBER("1")', '=ISTEXT(In!B13)', '=ISLOGICAL(1)', '=ISEVEN(-2.5)', '=ISODD(3.9)',
    '=ISERR(NA())', '=ISNA(NA())', '=ISNONTEXT(In!C3)', '=HEX2DEC("FF")', '=BIN2DEC("1111111111")', '=OCT2DEC("777")',
    # Fixed after the first run, and their neighbours
    '=CEILING(-2.5,2)', '=CEILING(2.5,-1)', '=FLOOR(-2.5,2)', '=FLOOR(2.5,0)', '=CEILING(2.5,0)', '=CEILING.MATH(-2.1,1,1)',
    '=FLOOR.MATH(-2.1,1,1)', '=CEILING.MATH(2.1,0.5)', '=FLOOR.MATH(-7,-2)', '=CEILING(-4.5,-2)', '=FLOOR(-4.5,-2)',
    '="1,234"+0', '="$1,234.5"+0', '="(5)"+0', '="1,23"+0', '="12,345,678"*1', '=VALUE("$5")', '=VALUE("(1,000)")',
    '=1E+15&""', '=1E+14&""', '=-1E+16&""', '=123456789012345&""', '=0.0000001&""', '=1.5E-10&""', '=2^60&""',
    '=MATCH(45,In!I1:I10,-1)', '=MATCH(100,In!I1:I10,-1)', '=MATCH(5,In!I1:I10,-1)', '=MATCH(101,In!I1:I10,-1)',
    '=XLOOKUP("c",In!F1:F10,In!E1:E10)', '=XLOOKUP(99,In!E1:E10,In!F1:F10)', '=XLOOKUP(95,In!E1:E10,In!F1:F10,"-",1)',
    '=XLOOKUP("D*",In!F1:F10,In!E1:E10,,2)', '=XLOOKUP(40,In!E1:E10,In!F1:F10,,0,-1)',
    '=TEXTJOIN(",",TRUE,In!C1:C6)', '=CONCAT(In!A1:A3)', '=IFS(In!A1>0,"pos",In!A1=0,"zero")', '=MAXIFS(In!A1:A15,In!A1:A15,"<100")',
    '=STDEV.S(In!E1:E10)', '=VAR.P(In!E1:E10)', '=RANK.EQ(30,In!E1:E10)', '=DAYS("2024-03-01","2024-02-01")',
    # Date and time text in arithmetic
    '="2024-03-01"-"2024-02-01"', '="15-Jan-2024"+0', '="Jan 15, 2024"+0', '="15 January 2024"+0', '="14:30"+0',
    '="2:30 PM"*24', '="2024-02-30"+0', '=DATEVALUE("15-Jan-2024")', '=DATEVALUE("2024-01-31 12:00")', '=VALUE("12:00")',
    '="2024-01-31 18:00"+0', '=YEAR("2024-07-04")', '=EDATE("2024-01-31",1)', '="12:00 AM"+0', '="13:00 PM"+0', '="24:00"+0',
    '="2023-02-29"+0', '="1900-02-29"+0', '=DATEVALUE("12:00")',
    # Text order: punctuation, digits, letters, accents
    '="~"<"b"', '="1"<"a"', '="-"<"1"', '="_"<"a"', '="["<"a"', '="é"="e"', '="é"<"f"', '="E"<"é"', '="a-b"<"ab"',
    '="Zebra">"apple"', '="a1"<"a10"', '="a2"<"a10"', '=" a"<"a"', '="ß"="ss"', '=MATCH("é",{"a","e","é","f"},0)',
    # Date criteria
    '=COUNTIF(In!A1:A15,">1900-01-10")', '=SUMIF(In!A1:A15,"<=1900-01-03")', '=COUNTIFS(In!A1:A15,">=2024-01-01",In!A1:A15,"<2024-12-31")',
    # Arrays
    '=SUM(In!E1:E3*In!G1:G3)', '=MDETERM(In!E1:F2)', '=SUMPRODUCT(--(In!F1:F10>"c"))', '=SUM((In!E1:E10>50)+0)',
]

NUM_REFS = [f'In!A{i}' for i in range(1, 16)]
TXT_REFS = [f'In!B{i}' for i in range(1, 16)]
MIX_REFS = [f'In!C{i}' for i in range(1, 16)]
RANGES = ['In!A1:A15', 'In!E1:E10', 'In!G1:G10', 'In!C1:C15', 'In!A1:A5']


def rnd_num(r):
    return r.choice(['0', '1', '2', '3', '-1', '0.5', '2.5', '10', '0.1', '1E+3', '7', '-2.5'])


def num_expr(r, d):
    if d <= 0 or r.random() < 0.3:
        return r.choice([rnd_num(r), r.choice(NUM_REFS), r.choice(NUM_REFS), r.choice(MIX_REFS)])
    k = r.randrange(16)
    a = lambda: num_expr(r, d - 1)  # noqa: E731
    if k == 0:
        return f'({a()}{r.choice(["+", "-", "*", "/"])}{a()})'
    if k == 1:
        return f'-{a()}'
    if k == 2:
        return f'ROUND({a()},{r.choice(["0", "1", "2", "-1"])})'
    if k == 3:
        return f'MOD({a()},{r.choice(["2", "3", "-3", "0.5"])})'
    if k == 4:
        return f'IF({bool_expr(r, d - 1)},{a()},{a()})'
    if k == 5:
        return f'{r.choice(["SUM", "AVERAGE", "MAX", "MIN", "COUNT", "COUNTA", "PRODUCT", "MEDIAN"])}({r.choice(RANGES)})'
    if k == 6:
        return f'{r.choice(["INT", "ABS", "TRUNC", "SIGN", "EVEN", "ODD"])}({a()})'
    if k == 7:
        return f'LEN({str_expr(r, d - 1)})'
    if k == 8:
        return f'COUNTIF({r.choice(RANGES)},"{r.choice([">", "<", ">=", "<=", "<>", "="])}{rnd_num(r)}")'
    if k == 9:
        return f'SUMIF({r.choice(RANGES[:3])},"{r.choice([">", "<", "<>"])}{rnd_num(r)}")'
    if k == 10:
        return f'IFERROR({a()},{rnd_num(r)})'
    if k == 11:
        return f'({a()}^{r.choice(["2", "0.5", "-1", "3"])})'
    if k == 12:
        return f'N({a()})'
    if k == 13:
        return f'MATCH({a()},In!E1:E10,{r.choice(["1", "0", "-1"])})'
    if k == 14:
        return f'INDEX(In!E1:E10,{r.choice(["1", "3", "10"])})'
    return f'{r.choice(["ROUNDUP", "ROUNDDOWN"])}({a()},{r.choice(["0", "1"])})'


def str_expr(r, d):
    if d <= 0 or r.random() < 0.35:
        return r.choice(['"abc"', '"A b"', '""', r.choice(TXT_REFS), r.choice(TXT_REFS), r.choice(MIX_REFS)])
    k = r.randrange(9)
    s = lambda: str_expr(r, d - 1)  # noqa: E731
    if k == 0:
        return f'({s()}&{s()})'
    if k == 1:
        return f'UPPER({s()})'
    if k == 2:
        return f'LEFT({s()},{r.choice(["1", "2", "0", "5"])})'
    if k == 3:
        return f'MID({s()},{r.choice(["1", "2", "3"])},{r.choice(["1", "2", "9"])})'
    if k == 4:
        return f'TRIM({s()})'
    if k == 5:
        return f'SUBSTITUTE({s()},"a","Z")'
    if k == 6:
        return f'({num_expr(r, d - 1)}&"")'
    if k == 7:
        return f'IF({bool_expr(r, d - 1)},{s()},{s()})'
    return f'REPT({s()},2)'


def bool_expr(r, d):
    if d <= 0 or r.random() < 0.3:
        return r.choice(['TRUE', 'FALSE', r.choice(MIX_REFS)])
    k = r.randrange(6)
    op = r.choice(['=', '<>', '<', '>', '<=', '>='])
    if k == 0:
        return f'({num_expr(r, d - 1)}{op}{num_expr(r, d - 1)})'
    if k == 1:
        return f'({str_expr(r, d - 1)}{op}{str_expr(r, d - 1)})'
    if k == 2:
        return f'{r.choice(["AND", "OR", "XOR"])}({bool_expr(r, d - 1)},{bool_expr(r, d - 1)})'
    if k == 3:
        return f'NOT({bool_expr(r, d - 1)})'
    if k == 4:
        return f'{r.choice(["ISNUMBER", "ISTEXT", "ISBLANK", "ISERROR", "ISLOGICAL"])}({r.choice(MIX_REFS + TXT_REFS)})'
    return f'({str_expr(r, d - 1)}{op}{num_expr(r, d - 1)})'


# Functions newer than Excel 2007 are written with a _xlfn. prefix in files; without it
# LibreOffice reads them as unknown names.
MODERN = ['XLOOKUP', 'IFS', 'SWITCH', 'XOR', 'IFNA', 'CONCAT', 'TEXTJOIN', 'DAYS', 'MAXIFS', 'MINIFS',
          'CEILING.MATH', 'FLOOR.MATH', 'STDEV.S', 'STDEV.P', 'VAR.S', 'VAR.P', 'RANK.EQ']


def xlfn(f):
    import re
    return re.sub(r'(?<![A-Za-z0-9_.])(' + '|'.join(re.escape(m) for m in MODERN) + r')\(', lambda m: f'_xlfn.{m.group(1)}(', f)


CURATED += [text_probe(v, c) for c in FORMAT_CODES for v in FORMAT_VALUES]


def main():
    r = random.Random(SEED)
    probes = list(CURATED)
    for _ in range(COUNT):
        kind = r.random()
        probes.append('=' + (num_expr(r, 3) if kind < 0.55 else str_expr(r, 3) if kind < 0.8 else bool_expr(r, 3)))

    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = 'In'
    for i in range(15):
        ins.cell(i + 1, 1, NUMS[i])
        if TEXT[i] != '':
            ins.cell(i + 1, 2, TEXT[i])
        if MIXED[i] is not None:
            ins.cell(i + 1, 3, MIXED[i])
    for i in range(10):
        ins.cell(i + 1, 5, SORTED[i])
        ins.cell(i + 1, 6, NAMES[i])
        ins.cell(i + 1, 7, UNSORTED[i])
    for i, v in enumerate([-100, 30, 40, 50, 20]):
        ins.cell(i + 1, 8, v)
    for i in range(10):  # descending, for MATCH(..., -1)
        ins.cell(i + 1, 9, SORTED[9 - i])
    for i in range(10):  # a row for HLOOKUP
        ins.cell(12, 5 + i, SORTED[i])
        ins.cell(13, 5 + i, NAMES[i])
    p = wb.create_sheet('P')
    for i, f in enumerate(probes):
        p.cell(i + 1, 1, xlfn(f))
    src = os.path.join(WORK, 'probe-src.xlsx')
    wb.save(src)

    doc = lo.open_doc(src)
    lo.recalc_save(doc, os.path.join(WORK, 'probe.xlsx'))
    lo.close(doc)
    lo.stop()
    print(f'{len(probes)} probes ({len(CURATED)} curated) → work/probe.xlsx')


if __name__ == '__main__':
    main()
